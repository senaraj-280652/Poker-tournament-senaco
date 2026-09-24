# -*- coding: utf-8 -*-
"""Journal DURABLE des actions effectuées depuis le contrôle à distance
par les utilisateurs ADMIN/DIRTO (chantier "LOG", Phase 1 — moteur de
journalisation uniquement, sans encore aucun branchement visuel dans
l'onglet CA/LOG).

Stocké dans ~/.poker_tournament/actions_log.sqlite3 (même mécanisme de
répertoire global que roster.py/license.py/open_windows.py), donc
GLOBAL à l'application — indépendant de chaque fichier .tournoi,
survit à la fermeture d'un tournoi, à la fermeture de l'application, à
un redémarrage, à un changement de tournoi.

Journal de TRANSACTIONS utiles uniquement (demande explicite) : ce
module ne décide JAMAIS lui-même ce qui mérite d'être journalisé — voir
remote_control.py pour la liste exacte des points d'appel (actions
mutantes ADMIN/DIRTO réellement issues du contrôle à distance, jamais
un sondage/une lecture passive).

Robustesse ABSOLUE (exigence explicite) : log_action() n'échoue JAMAIS
et ne bloque JAMAIS l'appelant — une panne du LOG (disque plein, fichier
verrouillé, permissions...) ne doit jamais empêcher une action réelle du
tournoi de s'exécuter. Même philosophie que remote_control.py:
log_remote_event (journal diagnostique existant, sans rapport avec
celui-ci) : "N'échoue JAMAIS".

Concurrence : plusieurs fenêtres/tournois sont des PROCESSUS séparés
(voir "Menu principal", main.py) pouvant écrire simultanément dans ce
même fichier global — d'où le mode WAL (conçu justement pour des
lecteurs/écrivains concurrents, y compris inter-processus), combiné à
une connexion SQLite JETABLE à chaque écriture (jamais une connexion
partagée entre threads ni entre appels : leçon déjà tirée dans ce
projet du crash cross-thread SQLite du contrôle à distance, voir
l'historique de remote_control.py) — ouverte, utilisée, refermée
immédiatement par log_action(), jamais conservée."""
import os
import sqlite3
import threading
import time
import unicodedata
from datetime import datetime

RESULT_SUCCESS = "SUCCESS"
RESULT_ACCEPTED = "ACCEPTED"
RESULT_DENIED = "DENIED"
RESULT_ERROR = "ERROR"
VALID_RESULTS = (RESULT_SUCCESS, RESULT_ACCEPTED, RESULT_DENIED, RESULT_ERROR)

# Micro-tentatives face à un "database is locked" transitoire — voir la
# docstring de log_action. Bornées et courtes : ne doivent jamais faire
# de ce journal secondaire un frein perceptible pour l'action réelle.
_MAX_ATTEMPTS = 3
_RETRY_DELAY_SECONDS = 0.05


def _log_dir():
    home = os.path.expanduser("~")
    d = os.path.join(home, ".poker_tournament")
    os.makedirs(d, exist_ok=True)
    return d


def _log_path():
    return os.path.join(_log_dir(), "actions_log.sqlite3")


# CHECK sur "result" : les 4 valeurs sont un contrat stable (voir
# VALID_RESULTS ci-dessus) — SQLite ne permet pas de modifier un CHECK
# existant par ALTER TABLE, mais une future 5e valeur resterait une
# migration simple (recréer la table, copier les lignes) précisément
# PARCE QUE le schéma reste par ailleurs minimal (pas d'autre contrainte
# forte) : ce CHECK n'ajoute donc pas de complexité de migration au-delà
# de celle, déjà nécessaire, d'ajouter une colonne.
_SCHEMA = """
CREATE TABLE IF NOT EXISTS actions_log (
    id              INTEGER PRIMARY KEY AUTOINCREMENT,
    ts              TEXT    NOT NULL,
    tournament_name TEXT    NOT NULL,
    tournament_path TEXT    NOT NULL,
    user_name       TEXT,
    role            TEXT    NOT NULL,
    device_id       TEXT,
    device_label    TEXT,
    category        TEXT    NOT NULL,
    action          TEXT    NOT NULL,
    player_name     TEXT,
    result          TEXT    NOT NULL CHECK (result IN ('SUCCESS', 'ACCEPTED', 'DENIED', 'ERROR')),
    message         TEXT
);

CREATE INDEX IF NOT EXISTS idx_actions_log_ts        ON actions_log(ts);
CREATE INDEX IF NOT EXISTS idx_actions_log_tournament ON actions_log(tournament_path);
CREATE INDEX IF NOT EXISTS idx_actions_log_user       ON actions_log(user_name);
CREATE INDEX IF NOT EXISTS idx_actions_log_category   ON actions_log(category);
CREATE INDEX IF NOT EXISTS idx_actions_log_player     ON actions_log(player_name);
"""


# _ensure_schema exécutait "CREATE TABLE/INDEX IF NOT EXISTS" à CHAQUE
# appel de log_action — idempotent, mais chaque instruction acquiert
# quand même brièvement le verrou d'écriture WAL. Sous forte
# concurrence (plusieurs threads/processus écrivant en rafale), ça
# multipliait inutilement les occasions de contention juste avant
# l'INSERT réel. `_schema_ready_paths` (mémoire du PROCESSUS courant
# uniquement, jamais persisté) fait que le schéma n'est réellement
# exécuté qu'UNE SEULE FOIS par processus et par fichier — sans danger
# pour deux processus démarrant en même temps : IF NOT EXISTS reste la
# garantie ultime, ceci n'est qu'une optimisation de contention.
_schema_ready_paths = set()
_schema_ready_lock = threading.Lock()


def _ensure_schema(conn, path):
    if path in _schema_ready_paths:
        return
    with _schema_ready_lock:
        if path in _schema_ready_paths:
            return
        conn.executescript(_SCHEMA)
        _schema_ready_paths.add(path)


def _open_connection(path):
    """Connexion SQLite jetable (voir docstring du module) : WAL +
    synchronous=NORMAL (couple standard recommandé avec WAL), timeout de
    5s pour absorber sans erreur une contention transitoire entre
    processus concurrents plutôt que de lever immédiatement "database is
    locked"."""
    conn = sqlite3.connect(path, timeout=5.0)
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA synchronous=NORMAL")
    return conn


def log_action(tournament_name, tournament_path, role, category, action, result,
                user_name=None, device_id=None, device_label=None,
                player_name=None, message=None):
    """Ajoute une ligne au journal — voir la docstring du module pour la
    garantie de robustesse : n'échoue JAMAIS, ne lève JAMAIS, quelle que
    soit la cause (disque plein, fichier verrouillé au-delà du timeout,
    permissions, résultat invalide...). `result` DOIT être l'une des 4
    valeurs de VALID_RESULTS ; toute autre valeur est silencieusement
    ramenée à RESULT_ERROR plutôt que de faire échouer l'insertion (le
    CHECK SQLite la refuserait sinon, ce qui romprait la garantie
    "jamais bloquant") — jamais une exception qui remonte à l'appelant,
    même dans ce cas limite.

    Quelques micro-tentatives (voir _MAX_ATTEMPTS ci-dessous) en cas de
    "database is locked" transitoire : observé en pratique (tests de
    charge, plusieurs threads/processus écrivant en rafale sans aucun
    délai réaliste entre deux actions) que le busy_timeout de la
    connexion (5s, voir _open_connection) n'empêche pas TOUJOURS cette
    erreur ponctuelle malgré un délai généreux — quelques retries
    quasi-immédiats suffisent largement en usage réel (des actions
    humaines, jamais deux processus écrivant à la microseconde près).
    Reste borné et silencieux : la dernière tentative échouée est, comme
    toute autre panne, avalée sans jamais bloquer l'appelant."""
    for attempt in range(_MAX_ATTEMPTS):
        try:
            if result not in VALID_RESULTS:
                result = RESULT_ERROR
            ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            path = _log_path()
            conn = _open_connection(path)
            try:
                _ensure_schema(conn, path)
                conn.execute(
                    "INSERT INTO actions_log "
                    "(ts, tournament_name, tournament_path, user_name, role, "
                    " device_id, device_label, category, action, player_name, "
                    " result, message) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    (
                        ts, tournament_name or "", tournament_path or "",
                        user_name, role or "", device_id, device_label,
                        category or "", action or "", player_name, result, message,
                    ),
                )
                conn.commit()
            finally:
                conn.close()
            return
        except sqlite3.OperationalError:
            if attempt + 1 >= _MAX_ATTEMPTS:
                return
            time.sleep(_RETRY_DELAY_SECONDS)
        except Exception:
            return


# =======================================================================
# API DE LECTURE (chantier "LOG", Phase 2, 2026-09-24) — vient s'ajouter
# à log_action() ci-dessus, qui reste ENTIÈREMENT inchangée (moteur
# d'écriture Phase 1). Objectif explicite : centraliser tout le SQL de
# lecture ICI, jamais dispersé dans main.py (voir App._refresh_log_tab/
# _refresh_log_filter_choices, qui n'appellent QUE ces fonctions).
#
# Robustesse : mêmes principes que log_action (connexion SQLite jetable,
# jamais partagée ; micro-tentatives bornées sur un verrou transitoire),
# mais PAS la même garantie absolue "jamais bloquant pour une action du
# tournoi" — une fonction de lecture n'est jamais appelée depuis le
# thread HTTP du contrôle à distance, seulement depuis l'onglet LOG côté
# Tk, en réponse à une action explicite de l'utilisateur (Rechercher,
# Réinitialiser, retour sur l'onglet). Une panne de lecture reste malgré
# tout totalement silencieuse (résultat vide plutôt qu'une exception) :
# consulter le journal ne doit jamais faire planter l'application.
# =======================================================================

# Catégories réellement utilisées par remote_control.py (voir sa
# docstring de _LOGGED_FIRE_AND_FORGET_ACTIONS et _ACTION_PERMISSION) —
# ORDRE volontairement fixe (jamais alphabétique) : reflète l'ordre
# naturel des fonctions du contrôle à distance tel qu'il apparaît déjà
# ailleurs dans l'application (voir database.REMOTE_PERMISSION_LABELS),
# pas l'ordre des chaînes techniques. list_categories() ci-dessous ne
# retourne QUE celles réellement présentes en base, dans cet ordre.
_CATEGORY_ORDER = ["eliminations", "moves", "clock", "levels", "photos", "rebalance", "admin_only"]

CATEGORY_LABELS = {
    "eliminations": "Éliminations",
    "moves": "Mouvements",
    "clock": "Chronomètre",
    "levels": "Niveaux",
    "photos": "Photos",
    "rebalance": "Rééquilibrage",
    "admin_only": "Administration",
}

# Les 12 valeurs d'"action" réellement écrites par remote_control.py —
# libellés français lisibles pour la colonne "Action" du tableau LOG
# (jamais les clés techniques brutes, demande explicite du 2026-09-24).
ACTION_LABELS = {
    "eliminate": "Élimination",
    "confirm_move": "Mouvement confirmé",
    "terminer": "Terminé (bouton)",
    "toggle_pause": "Pause / Reprise",
    "chronometre": "Affichage du chrono",
    "niveau_precedent": "Niveau précédent",
    "niveau_suivant": "Niveau suivant",
    "elimination": "Onglet Joueurs (⏸)",
    "end_tournament": "Fin du tournoi",
    "upload_photo": "Photo ajoutée",
    "delete_photo": "Photo supprimée",
    "rebalance_answer": "Réponse au rééquilibrage (UTG)",
}

# "NONE" (appareil approuvé mais jamais lié à un propriétaire) affiché
# "Non lié" (demande explicite du 2026-09-24) — ADMIN/DIRTO restent tels
# quels, déjà affichés bruts partout ailleurs dans l'application (voir
# roster.py, aucun libellé français n'existe pour ces deux groupes).
ROLE_LABELS = {"NONE": "Non lié"}

RESULT_LABELS = {
    RESULT_SUCCESS: "Réussi",
    RESULT_ACCEPTED: "Accepté",
    RESULT_DENIED: "Refusé",
    RESULT_ERROR: "Erreur",
}


def category_label(category):
    """Libellé français d'une catégorie — repli sur la valeur technique
    brute si jamais inconnue (défensif, ne doit normalement pas arriver :
    voir _CATEGORY_ORDER/CATEGORY_LABELS, tenus synchronisés avec
    remote_control.py)."""
    return CATEGORY_LABELS.get(category, category)


def action_label(action):
    """Libellé français d'une action — même repli défensif que
    category_label ci-dessus."""
    return ACTION_LABELS.get(action, action)


def role_label(role):
    """Libellé français d'un rôle — seul "NONE" a un libellé dédié
    (voir ROLE_LABELS) ; "ADMIN"/"DIRTO" (et toute valeur future
    inconnue) sont renvoyés tels quels."""
    return ROLE_LABELS.get(role, role)


def result_label(result):
    """Libellé français d'un résultat — repli défensif sur la valeur
    technique brute si jamais hors de VALID_RESULTS (ne devrait jamais
    arriver, le CHECK SQLite l'interdit)."""
    return RESULT_LABELS.get(result, result)


# Mêmes bornes que pour l'écriture (voir _MAX_ATTEMPTS/_RETRY_DELAY_
# SECONDS ci-dessus) — réutilisées telles quelles pour la lecture,
# même philosophie de robustesse ("quelques micro-tentatives quasi-
# immédiates suffisent largement en usage réel").
def _open_read_connection(path):
    """Connexion SQLite jetable EN LECTURE SEULE (query_only=ON — garde-
    fou défensif, cette famille de fonctions ne doit STRUCTURELLEMENT
    jamais pouvoir écrire). Pas de PRAGMA journal_mode ici : le mode WAL
    est une propriété PERSISTANTE du fichier lui-même (fixée par
    _open_connection dès la première écriture), une nouvelle connexion
    en lecture le voit déjà sans avoir besoin de le redéfinir — et une
    telle PRAGMA serait de toute façon incompatible avec query_only."""
    conn = sqlite3.connect(path, timeout=5.0)
    conn.execute("PRAGMA query_only = ON")
    conn.row_factory = sqlite3.Row
    return conn


def _read_query(sql, params=()):
    """Exécute `sql` sur une connexion de lecture jetable, avec les mêmes
    micro-tentatives que log_action en cas de verrou transitoire. Fichier
    absent : [] SANS jamais créer la base (ouvrir une connexion sqlite3
    sur un chemin inexistant la créerait sinon silencieusement — un
    effet de bord que la seule CONSULTATION du journal ne doit jamais
    produire). Toute autre panne : [] également, jamais d'exception qui
    remonterait à l'appelant (onglet LOG, main.py)."""
    path = _log_path()
    if not os.path.exists(path):
        return []
    for attempt in range(_MAX_ATTEMPTS):
        try:
            conn = _open_read_connection(path)
            try:
                return conn.execute(sql, params).fetchall()
            finally:
                conn.close()
        except sqlite3.OperationalError:
            if attempt + 1 >= _MAX_ATTEMPTS:
                return []
            time.sleep(_RETRY_DELAY_SECONDS)
        except Exception:
            return []
    return []


def _distinct_non_empty(column):
    """Valeurs distinctes non NULL/non vides d'une colonne texte, triées
    alphabétiquement. `column` est TOUJOURS un littéral interne fixe
    (jamais une entrée utilisateur, voir les appelants ci-dessous) —
    l'interpolation directe dans le SQL est donc sûre ici, aucune valeur
    externe n'atteint jamais cette chaîne."""
    sql = (
        f"SELECT DISTINCT {column} FROM actions_log "
        f"WHERE {column} IS NOT NULL AND TRIM({column}) != '' "
        f"ORDER BY {column}"
    )
    return [r[0] for r in _read_query(sql)]


def _natural_sort_key_fr(value):
    """Clé de tri insensible à la casse ET aux accents (NFKD, marques
    combinantes supprimées, puis casefold) — demande explicite du
    2026-09-24, "tri naturel français" pour le combobox Joueur. N'altère
    JAMAIS la valeur elle-même : uniquement une clé de comparaison,
    l'orthographe/les accents d'origine restent inchangés à l'affichage
    (voir list_players, qui trie mais renvoie les valeurs telles
    quelles)."""
    normalized = unicodedata.normalize("NFKD", value or "")
    stripped = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    return stripped.casefold()


def list_tournaments():
    """[{"tournament_name":..., "tournament_path":...}, ...] — une entrée
    par tournament_path DISTINCT (jamais dédupliqué par nom seul : deux
    tournois différents peuvent légitimement partager le même nom, voir
    App._build_log_tournament_label_maps pour la désambiguïsation
    d'AFFICHAGE, une préoccupation d'interface qui reste hors de ce
    module). Triée par nom puis par chemin pour un ordre stable et
    reproductible d'un appel à l'autre."""
    rows = _read_query(
        "SELECT DISTINCT tournament_name, tournament_path FROM actions_log "
        "ORDER BY tournament_name, tournament_path"
    )
    return [{"tournament_name": r[0], "tournament_path": r[1]} for r in rows]


def list_users():
    """Noms d'utilisateur distincts réellement présents dans le journal
    (jamais NULL — voir test_role_none_sans_proprietaire) — pour peupler
    le combobox Utilisateur. Tri alphabétique simple (contrairement à
    list_players ci-dessous, un tri "à la française" n'a pas été demandé
    ici)."""
    return _distinct_non_empty("user_name")


def list_players():
    """Noms de joueurs distincts, triés de façon insensible à la casse
    ET aux accents (voir _natural_sort_key_fr) — l'orthographe/les
    accents d'origine, eux, restent inchangés dans les valeurs
    retournées."""
    return sorted(_distinct_non_empty("player_name"), key=_natural_sort_key_fr)


def list_categories():
    """Catégories RÉELLEMENT présentes dans le journal (jamais les 7
    valeurs théoriques de _CATEGORY_ORDER si certaines ne sont encore
    jamais survenues — un journal tout neuf n'a par exemple jamais vu
    "photos"), dans l'ordre _CATEGORY_ORDER. Une catégorie future
    inconnue (5e résultat déjà anticipé par le schéma, voir _SCHEMA plus
    haut : ce n'est pas le même genre d'évolution, mais le principe
    "ne jamais perdre une valeur réelle" est le même) serait ajoutée à
    la fin, triée, plutôt que silencieusement ignorée."""
    present = set(_distinct_non_empty("category"))
    ordered = [c for c in _CATEGORY_ORDER if c in present]
    ordered += sorted(present - set(_CATEGORY_ORDER))
    return ordered


def _build_filters_clause(ts_from=None, ts_to=None, tournament_path=None, user_name=None,
                           category=None, player_name=None):
    """(where_sql, params) COMMUNS à search_actions/count_actions/
    purge_actions ci-dessous (source UNIQUE de cette construction —
    demande du 2026-09-25, "Purger doit respecter exactement les mêmes
    filtres que Rechercher, jamais une logique dupliquée qui pourrait
    diverger") : chaque paramètre à None = aucune restriction sur ce
    critère, tous combinables (ET logique). `where_sql` inclut déjà le
    mot-clé WHERE (chaîne vide si aucun filtre actif, ce qui vaut alors
    "toutes les lignes")."""
    clauses = []
    params = []
    if ts_from:
        clauses.append("ts >= ?")
        params.append(ts_from)
    if ts_to:
        clauses.append("ts <= ?")
        params.append(ts_to)
    if tournament_path:
        clauses.append("tournament_path = ?")
        params.append(tournament_path)
    if user_name:
        clauses.append("user_name = ?")
        params.append(user_name)
    if category:
        clauses.append("category = ?")
        params.append(category)
    if player_name:
        clauses.append("player_name = ?")
        params.append(player_name)
    where_sql = f"WHERE {' AND '.join(clauses)}" if clauses else ""
    return where_sql, params


def search_actions(ts_from=None, ts_to=None, tournament_path=None, user_name=None,
                    category=None, player_name=None, limit=500):
    """Recherche en lecture seule dans le journal, PLUS RÉCENT en premier
    (`ORDER BY ts DESC, id DESC` — id en départage pour deux actions à la
    même seconde). Chaque paramètre de filtre à None = aucune restriction
    sur ce critère ; tous combinables (ET logique) — voir _build_filters_
    clause, partagée avec count_actions/purge_actions.

    `ts_from`/`ts_to` sont déjà les bornes ENTIÈREMENT formées
    ("AAAA-MM-JJ HH:MM:SS", comparables lexicographiquement grâce au
    format zéro-paddé de `ts` — voir log_action). La conversion depuis
    une saisie utilisateur JJ/MM/AAAA, ainsi que sa validation (format,
    Du <= Au), restent la responsabilité de l'appelant : une
    préoccupation d'interface, pas de ce module.

    `limit=None` (revu le 2026-09-25, "CE QUI CORRESPOND AUX FILTRES =
    CE QUI EST AFFICHÉ = CE QUI EST EXPORTÉ = CE QUI PEUT ÊTRE PURGÉ") :
    AUCUNE limite — toutes les lignes correspondantes sont renvoyées,
    sans troncature ni pagination, `truncated` valant alors toujours
    False. C'est ce que l'onglet LOG (main.py: App._refresh_log_tab)
    utilise désormais : il n'y a plus de plafond artificiel à 500 lignes
    à cet endroit. Le paramètre `limit` (nombre) reste néanmoins
    disponible pour d'autres usages qui en ont réellement besoin (tests,
    scénarios de pagination) — jamais retiré de cette fonction, seulement
    devenu optionnel.

    Retourne (rows, truncated) : `rows` = liste de dicts (une entrée par
    colonne de la table, au plus `limit` si celui-ci est un nombre),
    `truncated=True` si STRICTEMENT plus de `limit` lignes
    correspondaient — détecté via un seul `LIMIT limit+1` (jamais un
    COUNT(*) séparé, qui doublerait le coût de chaque recherche).
    Fichier absent ou panne : ([], False), comme un journal vide —
    jamais d'exception."""
    where_sql, params = _build_filters_clause(
        ts_from, ts_to, tournament_path, user_name, category, player_name,
    )
    if limit is None:
        sql = f"SELECT * FROM actions_log {where_sql} ORDER BY ts DESC, id DESC"
        rows = [dict(r) for r in _read_query(sql, params)]
        return rows, False

    sql = f"SELECT * FROM actions_log {where_sql} ORDER BY ts DESC, id DESC LIMIT ?"
    params = list(params) + [int(limit) + 1]

    rows = [dict(r) for r in _read_query(sql, params)]
    truncated = len(rows) > limit
    return rows[:limit], truncated


def count_actions(ts_from=None, ts_to=None, tournament_path=None, user_name=None, category=None, player_name=None):
    """Nombre EXACT d'opérations correspondant aux mêmes critères que
    purge_actions ci-dessous et, plus généralement, à ceux de search_
    actions (voir _build_filters_clause, partagée par les trois
    fonctions) : ce comptage sert UNIQUEMENT à préparer la confirmation
    d'une purge (voir App._on_log_purge, qui appelle cette fonction PUIS
    purge_actions avec EXACTEMENT les mêmes arguments, calculés une
    seule fois — jamais recalculés séparément, pour garantir qu'ils
    restent identiques).

    Revu le 2026-09-25, PUIS de nouveau le même jour ("je veux exactement
    la même interprétation que pour l'affichage/recherche" pour Du/Au) :
    `ts_from`/`ts_to` sont désormais TOUS LES DEUX FACULTATIFS, EXACTEMENT
    comme pour search_actions — None de chaque côté = pas de borne sur ce
    côté ("Du" vide = depuis le tout premier enregistrement, "Au" vide =
    jusqu'au tout dernier, les deux vides = toute la période disponible).
    Aucune borne n'est jamais fabriquée artificiellement ici : une
    absence de filtre se traduit simplement par l'absence de la clause
    SQL correspondante (voir _build_filters_clause), jamais par un
    "01/01/1900" ou une date du jour codée en dur.

    Lecture SEULE (voir _read_query : connexion jetable, query_only=ON,
    mêmes micro-tentatives que les autres fonctions de lecture du LOG —
    aucun DELETE ici). Respecte Tournoi/Utilisateur/Fonction/Joueur
    EXACTEMENT comme purge_actions, pour que le nombre annoncé dans la
    confirmation corresponde toujours à ce que purge_actions supprimera
    réellement — jamais un écart entre "ce qui est affiché/annoncé" et
    "ce qui est supprimé".

    Fichier absent -> 0 (jamais créé pour autant, comme les autres
    fonctions de lecture)."""
    where_sql, params = _build_filters_clause(
        ts_from, ts_to, tournament_path, user_name, category, player_name,
    )
    rows = _read_query(f"SELECT COUNT(*) FROM actions_log {where_sql}", params)
    return rows[0][0] if rows else 0


def purge_actions(ts_from=None, ts_to=None, tournament_path=None, user_name=None, category=None, player_name=None):
    """Supprime DÉFINITIVEMENT les lignes du journal correspondant aux
    critères fournis — `ts_from`/`ts_to` (déjà entièrement formées
    "AAAA-MM-JJ HH:MM:SS", même convention que search_actions) PLUS les
    mêmes 4 filtres optionnels que Rechercher — Tournoi/Utilisateur/
    Fonction/Joueur (`tournament_path`/`user_name`/`category`/
    `player_name`) — voir _build_filters_clause, la même construction
    que search_actions/count_actions, jamais dupliquée.

    REVU le 2026-09-25, PUIS de nouveau le même jour ("je veux exactement
    la même interprétation que pour l'affichage/recherche" pour Du/Au) :
    LES 6 CRITÈRES SONT DÉSORMAIS TOUS FACULTATIFS, None = aucune
    restriction sur ce critère, EXACTEMENT comme search_actions — plus
    aucun `ValueError` pour une date manquante. "Du" vide = depuis le
    tout premier enregistrement du journal ; "Au" vide = jusqu'au tout
    dernier ; les deux vides = toute la période disponible. Aucune borne
    n'est jamais fabriquée artificiellement ici (pas de "01/01/1900" ni
    de date du jour codée en dur) : une date absente se traduit
    simplement par l'absence de la clause SQL correspondante. La
    validation/conversion depuis une saisie utilisateur JJ/MM/AAAA (format,
    Du <= Au quand les deux sont renseignées) reste la responsabilité de
    l'appelant : voir App._on_log_purge.

    Voir count_actions ci-dessus (mêmes arguments, appelée par
    App._on_log_purge AVANT cette fonction pour annoncer le nombre
    EXACT de lignes concernées dans la confirmation) : les deux
    fonctions doivent toujours rester rigoureusement cohérentes entre
    elles sur ce qu'elles considèrent comme "dans le périmètre".

    Ne supprime JAMAIS le fichier lui-même, ne touche JAMAIS au schéma
    ni aux index (un simple DELETE, jamais un DROP/CREATE), et ne lance
    JAMAIS VACUUM automatiquement (demande explicite) — l'espace libéré
    reste simplement réutilisable par SQLite pour de futures écritures.

    Le CRITÈRE DE SUPPRESSION est EXACTEMENT le même que celui de
    Rechercher (search_actions), sans la limite d'affichage (`limit`) :
    "ce que Rechercher afficherait en entier, dans sa totalité, est ce
    que Purger supprime" — jamais un sous-ensemble tronqué par la
    pagination de l'écran, jamais un filtre affiché sans effet réel
    (principe fondamental redemandé explicitement : "ce qui correspond
    aux filtres et est affiché = ce qui sera purgé"). Un appel SANS
    AUCUN filtre (les 6 paramètres à None) supprime donc légitimement
    la TOTALITÉ du journal — comportement assumé, jamais empêché ici :
    la confirmation affichée à l'utilisateur (voir App._confirm_log_
    purge) est le garde-fou, pas cette fonction.

    Concurrence : même connexion jetable + micro-tentatives que
    log_action (voir _open_connection/_MAX_ATTEMPTS/_RETRY_DELAY_
    SECONDS) — le mode WAL absorbe une écriture (log_action) concurrente
    à cette suppression sans corruption, chacune sa propre transaction
    courte.

    Retourne le nombre EXACT de lignes supprimées (0 si le fichier
    n'existe pas encore — RIEN À PURGER, jamais créé pour autant — ou si
    rien ne correspond aux critères). Contrairement à log_action, une
    panne PERSISTANTE ici (au-delà des micro-tentatives) est RELEVÉE À
    L'APPELANT (jamais avalée) : cette fonction n'est appelée que sur
    demande explicite et déjà confirmée de l'utilisateur (bouton
    "Purger"), jamais depuis le thread HTTP du contrôle à distance — une
    panne doit donc être signalée, pas laisser croire à tort qu'une
    purge a réussi."""
    path = _log_path()
    if not os.path.exists(path):
        return 0

    where_sql, params = _build_filters_clause(
        ts_from, ts_to, tournament_path, user_name, category, player_name,
    )
    for attempt in range(_MAX_ATTEMPTS):
        try:
            conn = _open_connection(path)
            try:
                _ensure_schema(conn, path)
                cursor = conn.execute(f"DELETE FROM actions_log {where_sql}", params)
                conn.commit()
                return cursor.rowcount
            finally:
                conn.close()
        except sqlite3.OperationalError:
            if attempt + 1 >= _MAX_ATTEMPTS:
                raise
            time.sleep(_RETRY_DELAY_SECONDS)


# =======================================================================
# EXPORT (fenêtre "Exporter le LOG", correction du 2026-09-24 — remplace
# l'export CSV direct implémenté un peu plus tôt le même jour) : reprend
# le mécanisme déjà en place pour "Exporter les primes" (voir database.
# py: PRIMES_COLUMNS/_selected_period_columns/export_primes_csv/xlsx/pdf,
# main.py: PrimesExportDialog) — CSV (";"/utf-8-sig), Excel (openpyxl,
# même style d'en-tête vert/blanc), PDF (fpdf2). Volontairement PAS de
# section "Tableau à exporter" (Primes en a deux — Récapitulatif/
# Historique — LOG n'en a qu'UN).
#
# Ces 3 fonctions ne font AUCUNE requête SQL — elles écrivent un fichier
# à partir de `rows`, un instantané FIGÉ des lignes DÉJÀ AFFICHÉES dans
# self.log_tree au moment où la fenêtre d'export s'est ouverte (voir
# main.py: App._on_log_export/LogExportDialog) : "ce que je vois dans le
# tableau LOG = ce qui peut être exporté" (demande explicite) — jamais
# une nouvelle recherche silencieuse. Depuis le 2026-09-25 ("CE QUI
# CORRESPOND AUX FILTRES = CE QUI EST AFFICHÉ = CE QUI EST EXPORTÉ = CE
# QUI PEUT ÊTRE PURGÉ"), l'onglet LOG n'applique plus AUCUN plafond de
# 500 lignes (voir search_actions(limit=None)) : `rows` peut donc
# légitimement contenir des centaines/milliers de lignes — ces 3
# fonctions n'ont jamais eu de limite propre, rien à changer ici pour
# ça. Chaque valeur de `rows` est déjà un TEXTE prêt à l'affichage
# (libellés français, dates JJ/MM/AAAA — voir main.py: App._populate_
# log_tree, seule source de ces chaînes), jamais une valeur brute à
# reformater ici.
#
# `title`/`criteria_line` (PDF depuis le 2026-09-24, étendu à CSV/Excel
# le 2026-09-25 — demande explicite : "les 3 formats doivent afficher
# les mêmes critères et le même nombre d'opérations", ce qui n'avait de
# sens QUE pour PDF au départ) : mêmes paramètres, même comportement
# dans les 3 fonctions — `criteria_line` absent/vide n'affiche RIEN
# (jamais une ligne "Critères : (aucun)" par défaut), `title` absent
# retombe sur "Journal des actions". La ligne "Nombre d'opérations : N"
# est en revanche TOUJOURS affichée (jamais conditionnelle), calculée
# comme `len(rows)` DANS la fonction d'export elle-même — jamais un
# compte transmis séparément par l'appelant, qui pourrait diverger de ce
# qui est réellement écrit dans le fichier.
# =======================================================================

# Colonnes exportables, dans l'ORDRE D'AFFICHAGE FIXE du tableau LOG
# (jamais l'ordre dans lequel l'utilisateur les a cochées, voir
# selected_log_export_columns ci-dessous — même principe que database.
# _selected_period_columns) — device_id/tournament_path n'y figurent
# JAMAIS (demande explicite, "données interdites").
LOG_EXPORT_COLUMNS = [
    ("ts", "Date/Heure"),
    ("tournament", "Tournoi"),
    ("user", "Utilisateur"),
    ("role", "Rôle"),
    ("category", "Fonction"),
    ("action", "Action"),
    ("player", "Joueur"),
    ("result", "Résultat"),
    ("message", "Message"),
]


def selected_log_export_columns(keys):
    """Sous-ensemble de LOG_EXPORT_COLUMNS correspondant à `keys` (les
    colonnes cochées dans LogExportDialog), dans l'ORDRE D'ORIGINE de
    LOG_EXPORT_COLUMNS — jamais l'ordre de `keys` lui-même. `keys=None`
    renvoie toutes les colonnes."""
    if keys is None:
        return list(LOG_EXPORT_COLUMNS)
    keys = set(keys)
    return [c for c in LOG_EXPORT_COLUMNS if c[0] in keys]


# Substitutions LISIBLES pour quelques caractères de "typographie
# intelligente" courants (tiret cadratin/demi-cadratin, points de
# suspension, guillemets/apostrophes courbes, symbole euro) — hors
# latin-1 mais dont l'équivalent ASCII est évident et déjà utilisé
# ailleurs dans ce module. Tout le RESTE (emoji, symboles techniques,
# alphabets non couverts...) passe par le filet de sécurité général de
# _pdf_text ci-dessous, jamais par une liste de cas particuliers à
# rallonger indéfiniment.
_PDF_TEXT_SUBSTITUTIONS = {
    "—": "-",    # — tiret cadratin
    "–": "-",    # – tiret demi-cadratin
    "…": "...",  # … points de suspension
    "‘": "'",    # ' apostrophe courbe gauche
    "’": "'",    # ' apostrophe courbe droite
    "“": '"',    # " guillemet courbe gauche
    "”": '"',    # " guillemet courbe droit
    "€": "EUR",  # € euro
}


def _pdf_text(value):
    """Convertit `value` en texte SÛR pour la police CŒUR "Helvetica" de
    fpdf2 (jeu de caractères latin-1 STRICT, voir pdf.core_fonts_
    encoding) — appliqué à TOUT texte exporté en PDF (tournoi,
    utilisateur, fonction, action, joueur, résultat, message...).

    Diagnostic du 2026-09-24 : SANS ce traitement, le moindre caractère
    hors latin-1 (symbole, emoji, ponctuation "intelligente" saisie
    depuis un clavier de téléphone, alphabet non couvert...) fait lever
    fpdf.errors.FPDFUnicodeEncodingException au moment du rendu — une
    exception qui remontait alors jusqu'à Tkinter, avalée sans rien
    afficher à l'utilisateur (voir main.py: LogExportDialog._do_export,
    désormais aussi corrigé pour ne plus jamais rester silencieux dans
    ce genre de cas). Repéré via ACTION_LABELS["elimination"] =
    "Onglet Joueurs (⏸)", mais la protection ci-dessous est
    STRUCTURELLEMENT générale : n'importe quel champ libre (message,
    nom de joueur/tournoi/utilisateur) peut en théorie contenir
    n'importe quel caractère Unicode, jamais prévisible à l'avance.

    Traitement en 2 temps :
    1. Substitutions lisibles pour les cas fréquents identifiés dans
       _PDF_TEXT_SUBSTITUTIONS (voir ci-dessus).
    2. Pour chaque caractère restant hors latin-1 : tente une
       décomposition Unicode (NFKD) et ne garde que les composantes de
       base non diacritiques déjà en latin-1 (ex : "ế" -> "e" — utile
       pour des noms saisis avec des accents hors français) ; si le
       caractère reste irréductible (emoji, symboles techniques...), le
       remplace par "?" plutôt que de laisser fpdf2 lever une
       exception. Ne tronque JAMAIS le texte : chaque caractère
       d'origine produit toujours au moins un caractère de sortie —
       jamais de perte silencieuse d'information au-delà de ce
       caractère lui-même."""
    if value is None:
        return ""
    text = str(value)
    for original, replacement in _PDF_TEXT_SUBSTITUTIONS.items():
        text = text.replace(original, replacement)

    out_chars = []
    for ch in text:
        if ord(ch) <= 0xFF:
            out_chars.append(ch)
            continue
        decomposed_ok = False
        for base_ch in unicodedata.normalize("NFKD", ch):
            if not unicodedata.combining(base_ch) and ord(base_ch) <= 0xFF:
                out_chars.append(base_ch)
                decomposed_ok = True
        if not decomposed_ok:
            out_chars.append("?")
    return "".join(out_chars)


def export_csv(path, columns, rows, title=None, criteria_line=None):
    """CSV — séparateur ";", encodage "utf-8-sig" (Excel français), une
    colonne par entrée de `columns` (voir selected_log_export_columns),
    dans cet ordre. `rows` : voir la docstring de section ci-dessus.

    Revu le 2026-09-25 : avant les en-têtes de colonnes, 3 lignes à une
    seule cellule — le titre ("Journal des actions" par défaut), la
    ligne de critères SEULEMENT si `criteria_line` est non vide (jamais
    une ligne par défaut inventée), puis "Nombre d'opérations : N" —
    TOUJOURS présente, N = len(rows) (calculé ici, jamais transmis)."""
    import csv

    keys = [k for k, _ in columns]
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow([title or "Journal des actions"])
        if criteria_line:
            writer.writerow([criteria_line])
        writer.writerow([f"Nombre d'opérations : {len(rows)}"])
        writer.writerow([h for _, h in columns])
        for row in rows:
            writer.writerow([row.get(k, "") for k in keys])
    return path


def export_xlsx(path, columns, rows, title=None, criteria_line=None):
    """Excel (.xlsx) — même style que database.export_primes_xlsx (titre
    en gras, ligne d'en-tête verte/blanche, largeurs de colonnes
    ajustées) : reprend le mécanisme existant plutôt que d'en inventer un
    nouveau. `wrap_text` activé sur chaque cellule de donnée (ajout par
    rapport à Primes, aucune de ses colonnes n'a un texte aussi long que
    "Message" ici) pour rester lisible sans élargir démesurément la
    colonne. Nécessite 'openpyxl'.

    Revu le 2026-09-25 : sous le titre, `criteria_line` SEULEMENT si non
    vide (même convention que export_pdf, jamais une ligne par défaut
    inventée), puis "Nombre d'opérations : N" — TOUJOURS présente, N =
    len(rows). `header_row` reste calculé dynamiquement (ws.max_row + 1,
    inchangé) : ces lignes en plus n'imposent aucun numéro de ligne fixe
    ailleurs dans cette fonction."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    keys = [k for k, _ in columns]
    headers = [h for _, h in columns]

    wb = Workbook()
    ws = wb.active
    ws.title = "LOG"

    ws.append([title or "Journal des actions"])
    ws["A1"].font = Font(bold=True, size=14)
    if criteria_line:
        ws.append([criteria_line])
    ws.append([f"Nombre d'opérations : {len(rows)}"])
    ws.append([])

    header_row = ws.max_row + 1
    ws.append(headers)
    header_fill = PatternFill(start_color="1F4E24", end_color="1F4E24", fill_type="solid")
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        ws.append([row.get(k, "") for k in keys])
        for col in range(1, len(headers) + 1):
            ws.cell(row=ws.max_row, column=col).alignment = Alignment(vertical="top", wrap_text=True)

    widths = {
        "ts": 18, "tournament": 24, "user": 14, "role": 10,
        "category": 14, "action": 24, "player": 16, "result": 12, "message": 48,
    }
    for i, key in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(key, 16)

    wb.save(path)
    return path


def _pdf_wrap_row(pdf, x_start, y_start, col_widths, texts, row_height, bold, fill_rgb, text_rgb, align):
    """Dessine UNE ligne d'un tableau PDF fpdf2, chaque cellule pouvant
    s'étaler sur PLUSIEURS lignes (retour à la ligne automatique) tout en
    gardant toutes les cellules de cette ligne à la MÊME hauteur —
    fpdf2 n'offre pas nativement de tableau à hauteur de ligne variable
    avec cellules alignées (contrairement à database._write_pdf_table,
    qui suppose des valeurs courtes tenant sur une seule ligne — jamais
    le cas ici pour "Message", d'où cette variante dédiée). Mesure
    d'abord (multi_cell(..., dry_run=True, output="LINES"), sans
    dessiner) puis dessine chaque cellule (fond, bordure, texte) à la
    hauteur commune. Renvoie le y de fin de ligne, pour la ligne
    suivante."""
    pdf.set_font("Helvetica", "B" if bold else "", 9 if bold else 8)
    line_counts = [
        max(1, len(pdf.multi_cell(w, row_height, _pdf_text(t), dry_run=True, output="LINES")))
        for w, t in zip(col_widths, texts)
    ]
    cell_h = row_height * max(line_counts)

    x = x_start
    for w, t in zip(col_widths, texts):
        pdf.set_fill_color(*fill_rgb)
        pdf.rect(x, y_start, w, cell_h, style="F")
        pdf.set_draw_color(180, 180, 180)
        pdf.rect(x, y_start, w, cell_h, style="D")
        pdf.set_xy(x, y_start)
        pdf.set_text_color(*text_rgb)
        pdf.multi_cell(w, row_height, _pdf_text(t), border=0, align=align)
        x += w
    return y_start + cell_h


def format_log_export_criteria(date_from=None, date_to=None, tournament_label="Tous",
                                user_label="Tous", function_label="Toutes", player_label="Tous"):
    """Construit la ligne "Critères : ..." affichée sous le titre du PDF
    du LOG (demande du 2026-09-24, "petite amélioration du PDF
    uniquement") — App._on_log_export appelle cette fonction avec les
    valeurs EXACTEMENT telles qu'affichées dans les widgets de filtre
    au moment de l'export (aucun recalcul/nouvelle recherche ici, même
    principe que pour `rows`). `date_from`/`date_to` : déjà au format
    JJ/MM/AAAA affiché (ou None/"" si le champ est vide) ; les 4 autres
    labels sont déjà les valeurs de combobox affichées ("Tous"/"Toutes"
    par défaut, ou le libellé réellement sélectionné).

    Si les 6 critères sont TOUS à leur valeur par défaut (aucun filtre
    actif), renvoie simplement "Critères : Toutes les actions" — sinon
    une ligne listant Du/Au (uniquement si au moins l'un des deux est
    renseigné) puis TOUJOURS Tournoi/Utilisateur/Fonction/Joueur (même
    à leur valeur par défaut, voir l'exemple de la demande), séparés
    par " — ".

    Le résultat est un texte ordinaire — PAS encore passé par _pdf_text
    ici (cette fonction ne sait rien de fpdf2) : c'est export_pdf,
    seule à en avoir besoin, qui applique la protection Unicode PDF au
    moment du rendu, exactement comme pour tout autre texte exporté."""
    date_from = date_from or None
    date_to = date_to or None
    all_default = (
        date_from is None and date_to is None
        and tournament_label == "Tous" and user_label == "Tous"
        and function_label == "Toutes" and player_label == "Tous"
    )
    if all_default:
        return "Critères : Toutes les actions"

    segments = []
    if date_from and date_to:
        segments.append(f"Du {date_from} au {date_to}")
    elif date_from:
        segments.append(f"Du {date_from}")
    elif date_to:
        segments.append(f"Jusqu'au {date_to}")
    segments.append(f"Tournoi : {tournament_label}")
    segments.append(f"Utilisateur : {user_label}")
    segments.append(f"Fonction : {function_label}")
    segments.append(f"Joueur : {player_label}")
    return "Critères : " + " — ".join(segments)


def export_pdf(path, columns, rows, title=None, criteria_line=None):
    """PDF — même esprit que database._write_pdf_table (titre en gras,
    en-tête vert/blanc, lignes alternées, paysage automatique au-delà
    d'un certain nombre de colonnes, taille de police adaptée) mais avec
    un renfort explicite demandé pour LOG : chaque cellule (pas
    seulement "Message") peut s'étaler sur plusieurs lignes plutôt que
    de voir sa police rétrécie jusqu'à devenir illisible ou de déborder
    silencieusement — "ne pas tronquer silencieusement des informations
    importantes" (demande explicite). "Message" reçoit une largeur de
    colonne 2,2x supérieure aux autres (texte libre, souvent le plus
    long). Repagine automatiquement (saut de page + en-tête redessiné)
    si une ligne ne tient plus dans l'espace restant.

    `criteria_line` (ajouté le 2026-09-24, désormais aussi transmis à
    export_csv/export_xlsx depuis le 2026-09-25) : ligne optionnelle
    affichée en gris sous le titre, résumant les filtres Du/Au/Tournoi/
    Utilisateur/Fonction/Joueur actifs au moment de l'export — voir
    format_log_export_criteria. Rien n'est dessiné si absent/vide.
    Passe par _pdf_text comme tout le reste : un caractère hors latin-1
    dans un libellé de filtre (nom de tournoi, de joueur...) ne doit
    jamais faire échouer l'export, même protection que pour les
    colonnes du tableau.

    "Nombre d'opérations : N" (ajouté le 2026-09-25, N = len(rows)) :
    dessinée juste sous la ligne de critères (ou juste sous le titre si
    absente), même style gris — TOUJOURS présente, contrairement à
    `criteria_line`, et elle aussi protégée par _pdf_text (un nombre
    reste toujours latin-1, mais le principe "tout texte dessiné passe
    par _pdf_text" reste appliqué sans exception). Nécessite 'fpdf2'."""
    from fpdf import FPDF

    keys = [k for k, _ in columns]
    headers = [h for _, h in columns]
    orientation = "L" if len(columns) > 4 else "P"
    pdf = FPDF(orientation=orientation, unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, _pdf_text(title or "Journal des actions"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(90, 90, 90)
    if criteria_line:
        pdf.cell(0, 6, _pdf_text(criteria_line), new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 6, _pdf_text(f"Nombre d'opérations : {len(rows)}"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0, 0, 0)
    pdf.ln(4)

    avail_width = pdf.w - pdf.l_margin - pdf.r_margin
    weights = [2.2 if k == "message" else 1.0 for k in keys]
    total_weight = sum(weights)
    col_widths = [avail_width * w / total_weight for w in weights]
    row_height = 5

    def _draw_header():
        y = _pdf_wrap_row(
            pdf, pdf.l_margin, pdf.get_y(), col_widths, headers, row_height,
            bold=True, fill_rgb=(31, 78, 36), text_rgb=(255, 255, 255), align="C",
        )
        pdf.set_y(y)

    _draw_header()

    pdf.set_text_color(0, 0, 0)
    fill_toggle = False
    for row in rows:
        texts = [row.get(k, "") for k in keys]
        pdf.set_font("Helvetica", "", 8)
        needed_h = row_height * max(
            max(1, len(pdf.multi_cell(w, row_height, _pdf_text(t), dry_run=True, output="LINES")))
            for w, t in zip(col_widths, texts)
        )
        if pdf.get_y() + needed_h > pdf.page_break_trigger:
            pdf.add_page()
            _draw_header()

        fill_rgb = (247, 241, 227) if fill_toggle else (255, 255, 255)
        y = _pdf_wrap_row(
            pdf, pdf.l_margin, pdf.get_y(), col_widths, texts, row_height,
            bold=False, fill_rgb=fill_rgb, text_rgb=(0, 0, 0), align="L",
        )
        pdf.set_y(y)
        fill_toggle = not fill_toggle

    pdf.output(path)
    return path
