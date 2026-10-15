# -*- coding: utf-8 -*-
"""Groupe TMP — invités temporaires (demande du 2026-10-15) : participent
normalement au tournoi (table, équilibrages, éliminations) mais sont
exclus du CLASSEMENT OFFICIEL, des points (classement/présence/
assiduité), des bounties, et des statistiques officielles de
participation du club — tout en restant TRAÇABLES (éliminations,
historique) et en conservant une trace de leur passage en fin de
résultats.

AUCUNE nouvelle colonne : le champ `players.club` (déjà existant, déjà
figé PAR TOURNOI dès l'inscription, voir Database.add_player — jamais
resynchronisé avec le Répertoire ensuite) est le seul signal. Un joueur
est "TMP" si et seulement si son club, pour CE tournoi, correspond
EXACTEMENT (espaces/casse ignorés) à "TMP" — voir database.is_tmp_club.
Confirmé explicitement par l'utilisateur : aucun club nommé "TMP" n'a
jamais existé dans l'historique des tournois déjà enregistrés —> aucun
ancien tournoi n'est donc affecté par ce correctif (zéro migration,
calcul purement dérivé à la lecture + deux écritures ciblées, toujours
vers l'avenir).

Distinction centrale (demande explicite de l'utilisateur) :
- "effectif RÉEL" du tournoi (TMP compris) : onglet Joueurs ("rang"),
  "entries"/tapis moyen/effectifs du Lobby (Database.get_stats) — JAMAIS
  touchés par ce correctif.
- "effectif OFFICIEL" (classement/points/récompenses, TMP exclus) :
  onglet Classement ("official_rang"), 4 fonctions get_*_bonuses,
  build_period_summary (wins/best_place/tournaments_played/winner).

HORS PÉRIURÈTRE DÉLIBÉRÉ de ce correctif (pas couvert par les règles
explicites de l'utilisateur, qui ne parlent que de POINTS) : les gains
en argent réel (Database._results_rows/get_payouts_amounts, onglet
"Résultats"/"Gains" et leurs exports dédiés) restent basés sur l'ordre
RÉEL d'élimination, INCHANGÉS — voir ResultsRowsScopeBoundaryTest
ci-dessous, qui verrouille explicitement cette frontière de périmètre.

Utilise une vraie Database SQLite en fichier temporaire (même convention
que tests/test_ranking_formula.py) — jamais de doublure pour la logique
métier elle-même."""
import ast
import os
import shutil
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import database  # noqa: E402


def _main_source():
    main_py = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "main.py")
    with open(main_py, encoding="utf-8") as f:
        return f.read()


# =======================================================================
# 1. Fonctions pures (is_tmp_club / official_standings)
# =======================================================================
class IsTmpClubTest(unittest.TestCase):
    def test_correspondance_exacte_insensible_casse_et_espaces(self):
        for value in ("TMP", "tmp", "Tmp", "  TMP  ", "\tTMP\n"):
            self.assertTrue(database.is_tmp_club(value), repr(value))

    def test_jamais_une_sous_chaine(self):
        for value in ("TMP Poker Club", "Club TMP", "TMPX", "XTMP", "TMP2"):
            self.assertFalse(database.is_tmp_club(value), repr(value))

    def test_valeurs_vides_ou_absentes(self):
        self.assertFalse(database.is_tmp_club(""))
        self.assertFalse(database.is_tmp_club(None))
        self.assertFalse(database.is_tmp_club("   "))

    def test_club_regulier_quelconque(self):
        for value in ("CPC", "Autre Club", "", "Club du coin"):
            self.assertFalse(database.is_tmp_club(value))


class OfficialStandingsTest(unittest.TestCase):
    def test_exemple_exact_de_lutilisateur(self):
        """Pierre(1er, CPC), Visiteur(2e, TMP), Jacques(3e, CPC),
        Michel(4e, CPC) -> officiel : Pierre=1, Jacques=2, Michel=3,
        Visiteur hors classement (None)."""
        rows = [
            {"id": 1, "club": "CPC", "place": 1},
            {"id": 2, "club": "TMP", "place": 2},
            {"id": 3, "club": "CPC", "place": 3},
            {"id": 4, "club": "CPC", "place": 4},
        ]
        self.assertEqual(
            database.official_standings(rows), {1: 1, 2: None, 3: 2, 4: 3},
        )

    def test_tmp_vainqueur_reel_regulier_le_mieux_classe_devient_officiel(self):
        rows = [
            {"id": 10, "club": "TMP", "place": 1},
            {"id": 11, "club": "CPC", "place": 2},
            {"id": 12, "club": "CPC", "place": 3},
        ]
        self.assertEqual(
            database.official_standings(rows), {10: None, 11: 1, 12: 2},
        )

    def test_plusieurs_tmp_entremeles(self):
        rows = [
            {"id": 1, "club": "CPC", "place": 1},
            {"id": 2, "club": "TMP", "place": 2},
            {"id": 3, "club": "TMP", "place": 3},
            {"id": 4, "club": "CPC", "place": 4},
            {"id": 5, "club": "TMP", "place": 5},
            {"id": 6, "club": "CPC", "place": 6},
        ]
        self.assertEqual(
            database.official_standings(rows),
            {1: 1, 2: None, 3: None, 4: 2, 5: None, 6: 3},
        )

    def test_aucun_tmp_comportement_identique_a_lordre_brut(self):
        rows = [{"id": i, "club": "CPC", "place": i} for i in range(1, 6)]
        self.assertEqual(
            database.official_standings(rows), {i: i for i in range(1, 6)},
        )

    def test_tous_tmp_personne_de_classe(self):
        rows = [{"id": 1, "club": "TMP", "place": 1}, {"id": 2, "club": "TMP", "place": 2}]
        self.assertEqual(database.official_standings(rows), {1: None, 2: None})

    def test_place_none_jamais_classee(self):
        rows = [{"id": 1, "club": "CPC", "place": None}]
        self.assertEqual(database.official_standings(rows), {1: None})

    def test_ne_modifie_jamais_les_dicts_dorigine(self):
        rows = [{"id": 1, "club": "CPC", "place": 1}]
        database.official_standings(rows)
        self.assertEqual(rows, [{"id": 1, "club": "CPC", "place": 1}])

    def test_trou_sans_rapport_avec_un_tmp_najamais_promu_a_1(self):
        """PIÈGE RÉEL détecté en test (première version de cette fonction,
        un simple "renumérotage 1..M") : un SEUL joueur classé, place
        brute 2 (personne à la place 1 — ex. l'autre joueur a forfait
        AVANT toute élimination, voir Database.withdraw_player) ne doit
        JAMAIS devenir "rang officiel 1" juste parce qu'il est le seul de
        la liste — aucun TMP n'est impliqué ici, rien ne justifie de le
        faire "remonter". Seul un TMP de place brute STRICTEMENT
        inférieure peut faire remonter un régulier."""
        rows = [{"id": 1, "club": "CPC", "place": 2}]
        self.assertEqual(database.official_standings(rows), {1: 2})

    def test_trou_avec_tmp_ne_decale_que_du_nombre_de_tmp_strictement_avant(self):
        rows = [
            {"id": 1, "club": "TMP", "place": 1},
            {"id": 2, "club": "CPC", "place": 3},  # place 2 jamais attribuée (ex. forfait)
        ]
        # Un seul TMP de place strictement inférieure (1 < 3) -> décalage
        # de 1, jamais un renumérotage "1..M" qui donnerait à tort 1.
        self.assertEqual(database.official_standings(rows), {1: None, 2: 2})


# =======================================================================
# 2. Écritures : bounty jamais placée/transférée sur/depuis un TMP
# =======================================================================
class TmpBountyWriteSideTest(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmpdir, ignore_errors=True)
        self.db = database.Database(os.path.join(self.tmpdir, "A.tournoi"))
        self.db.set_setting("primes_enabled", 1)
        self.db.set_setting("bounty_amount", 100)
        self.addCleanup(self.db.close)

    def test_aucune_bounty_stampee_a_linscription_dun_tmp(self):
        pid = self.db.add_player("Visiteur", "TMP")
        self.assertEqual(self.db.get_player(pid)["bounty"], 0)

    def test_bounty_normale_pour_un_regulier(self):
        pid = self.db.add_player("Pierre", "CPC")
        self.assertEqual(self.db.get_player(pid)["bounty"], 100)

    def test_rebuy_dun_tmp_najoute_jamais_de_bounty(self):
        pid = self.db.add_player("Visiteur", "TMP")
        self.db.rebuy_player(pid)
        self.assertEqual(self.db.get_player(pid)["bounty"], 0)

    def test_rebuy_dun_regulier_reste_inchange(self):
        pid = self.db.add_player("Pierre", "CPC")
        self.db.rebuy_player(pid)
        self.assertEqual(self.db.get_player(pid)["bounty"], 200)

    def test_set_player_club_vers_tmp_remet_la_bounty_a_zero(self):
        pid = self.db.add_player("Michel", "CPC")
        self.assertEqual(self.db.get_player(pid)["bounty"], 100)
        self.db.set_player_club(pid, "TMP")
        self.assertEqual(self.db.get_player(pid)["bounty"], 0)

    def test_set_player_club_vers_regulier_ne_touche_pas_la_bounty(self):
        pid = self.db.add_player("Michel", "CPC")
        self.db.set_player_club(pid, "Autre Club")
        self.assertEqual(self.db.get_player(pid)["bounty"], 100)

    def test_regulier_eliminant_tmp_ne_recoit_rien(self):
        regular = self.db.add_player("Jacques", "CPC")
        tmp = self.db.add_player("Visiteur", "TMP")
        self.db.eliminate_player(tmp, eliminated_by_id=regular)
        jacques = self.db.get_player(regular)
        self.assertEqual(jacques["kills"], 0)
        self.assertEqual(jacques["bounty_won"], 0)

    def test_tmp_eliminant_regulier_ne_recoit_rien_et_bounty_abandonnee(self):
        tmp = self.db.add_player("Visiteur", "TMP")
        regular = self.db.add_player("Pierre", "CPC")  # porte une bounty de 100
        self.db.eliminate_player(regular, eliminated_by_id=tmp)
        visiteur = self.db.get_player(tmp)
        pierre = self.db.get_player(regular)
        self.assertEqual(visiteur["kills"], 0)
        self.assertEqual(visiteur["bounty_won"], 0)
        self.assertEqual(pierre["bounty"], 0)  # abandonnée, jamais transférée

    def test_tmp_eliminant_tmp_ne_recoit_rien(self):
        tmp1 = self.db.add_player("Visiteur1", "TMP")
        tmp2 = self.db.add_player("Visiteur2", "TMP")
        self.db.eliminate_player(tmp2, eliminated_by_id=tmp1)
        v1 = self.db.get_player(tmp1)
        self.assertEqual(v1["kills"], 0)
        self.assertEqual(v1["bounty_won"], 0)

    def test_regulier_eliminant_regulier_reste_inchange_non_regression(self):
        a = self.db.add_player("A", "CPC")
        b = self.db.add_player("B", "CPC")
        self.db.eliminate_player(b, eliminated_by_id=a)
        row_a = self.db.get_player(a)
        self.assertEqual(row_a["kills"], 1)
        self.assertEqual(row_a["bounty_won"], 100)

    def test_aucune_ligne_bounty_events_pour_une_elimination_impliquant_un_tmp(self):
        """Demande explicite : "conserver la trace des éliminations...
        sans attribution de récompense ni incrémentation des compteurs"
        — jamais de ligne bounty_events (aucun transfert n'a eu lieu),
        l'élimination elle-même reste traçable via eliminated_by_name/
        elim_time/elim_round (voir test suivant)."""
        tmp = self.db.add_player("Visiteur", "TMP")
        regular = self.db.add_player("Pierre", "CPC")
        before = self.db.conn.execute("SELECT COUNT(*) c FROM bounty_events").fetchone()["c"]
        self.db.eliminate_player(regular, eliminated_by_id=tmp)
        after = self.db.conn.execute("SELECT COUNT(*) c FROM bounty_events").fetchone()["c"]
        self.assertEqual(before, after)

    def test_elimination_dun_tmp_reste_tracable_dans_lhistorique(self):
        """L'élimination elle-même (qui, quand, par qui) n'est JAMAIS
        affectée par le groupe TMP — seule la prime en est absente."""
        regular = self.db.add_player("Jacques", "CPC")
        tmp = self.db.add_player("Visiteur", "TMP")
        self.db.eliminate_player(tmp, eliminated_by_id=regular)
        visiteur = self.db.get_player(tmp)
        self.assertEqual(visiteur["status"], "eliminated")
        self.assertEqual(visiteur["eliminated_by_name"], "Jacques")
        self.assertIsNotNone(visiteur["elim_time"])
        self.assertEqual(visiteur["place"], 2)  # 2 actifs avant cette élimination


# =======================================================================
# 3. Lectures : présence / assiduité / classement / bounty
# =======================================================================
class TmpReadSideTest(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmpdir, ignore_errors=True)
        self.db = database.Database(os.path.join(self.tmpdir, "A.tournoi"))
        self.db.set_setting("primes_enabled", 1)
        self.db.set_setting("attendance_bonus_points", 10)
        self.addCleanup(self.db.close)

    def test_aucun_point_de_presence_pour_un_tmp(self):
        self.db.add_player("Pierre", "CPC")
        self.db.add_player("Visiteur", "TMP")
        presence = self.db.get_presence_bonuses()
        self.assertEqual(presence["Pierre"], 10)
        self.assertEqual(presence.get("Visiteur", 0), 0)

    def test_aucun_point_dassiduite_pour_un_tmp(self):
        self.db.set_setting("assiduity_bonus_points", 20)
        self.db.set_setting("assiduity_consecutive_days", 1)
        self.db.add_player("Pierre", "CPC")
        self.db.add_player("Visiteur", "TMP")
        assiduite = {r["name"]: r for r in self.db.get_assiduity_bonuses()}
        self.assertEqual(assiduite["Pierre"]["points"], 20)
        self.assertEqual(assiduite["Visiteur"]["points"], 0)

    def test_classement_exemple_exact_de_lutilisateur(self):
        self.db.set_setting("ranking_formula", database.RANKING_FORMULA_CURRENT)
        pierre = self.db.add_player("Pierre", "CPC")
        visiteur = self.db.add_player("Visiteur", "TMP")
        jacques = self.db.add_player("Jacques", "CPC")
        michel = self.db.add_player("Michel", "CPC")
        # Éliminations (le premier éliminé prend la pire place) :
        # Michel(4e), Jacques(3e), Visiteur(2e), Pierre reste (1er).
        self.db.eliminate_player(michel, eliminated_by_id=jacques)
        self.db.eliminate_player(jacques, eliminated_by_id=visiteur)
        self.db.eliminate_player(visiteur, eliminated_by_id=pierre)

        ranking = {r["name"]: r["place"] for r in self.db.get_ranking_bonuses()}
        self.assertEqual(ranking, {"Pierre": 1, "Jacques": 2, "Michel": 3})
        self.assertNotIn("Visiteur", ranking)

    def test_tmp_vainqueur_reel_bascule_le_vainqueur_officiel(self):
        self.db.set_setting("ranking_formula", database.RANKING_FORMULA_CURRENT)
        pierre = self.db.add_player("Pierre", "CPC")
        visiteur = self.db.add_player("Visiteur", "TMP")
        jacques = self.db.add_player("Jacques", "CPC")
        self.db.eliminate_player(jacques, eliminated_by_id=pierre)
        self.db.eliminate_player(pierre, eliminated_by_id=visiteur)
        # Visiteur (TMP) reste actif -> vainqueur RÉEL.
        ranking = {r["name"]: r["place"] for r in self.db.get_ranking_bonuses()}
        self.assertEqual(ranking, {"Pierre": 1, "Jacques": 2})
        self.assertNotIn("Visiteur", ranking)

    def test_n_players_formule_exclut_les_tmp(self):
        """"Les TMP sont exclus du nombre de participants utilisé dans
        les formules de points" — vérifié en comparant la valeur réelle
        à database.ranking_points_table(n_regulier, ...)."""
        self.db.set_setting("ranking_formula", database.RANKING_FORMULA_CURRENT)
        a = self.db.add_player("A", "CPC")
        tmp1 = self.db.add_player("TMP1", "TMP")
        tmp2 = self.db.add_player("TMP2", "TMP")
        b = self.db.add_player("B", "CPC")
        # Élimine tout le monde sauf A (tournoi réellement terminé,
        # TMP compris, pour que "A" devienne vainqueur) : B(4e), TMP2(3e),
        # TMP1(2e), A reste (1er).
        self.db.eliminate_player(b, eliminated_by_id=a)
        self.db.eliminate_player(tmp2, eliminated_by_id=a)
        self.db.eliminate_player(tmp1, eliminated_by_id=a)
        ranking = {r["name"]: r for r in self.db.get_ranking_bonuses()}
        # 2 réguliers seulement (A, B) -> N=2 pour la formule, jamais 4.
        expected_valeur_vainqueur = database.ranking_points(1, 2, database.RANKING_FORMULA_CURRENT)
        self.assertEqual(ranking["A"]["valeur"], expected_valeur_vainqueur)
        self.assertEqual(ranking["A"]["place"], 1)
        self.assertEqual(ranking["B"]["place"], 2)
        self.assertNotIn("TMP1", ranking)
        self.assertNotIn("TMP2", ranking)

    def test_bounty_bonuses_exclut_entierement_les_tmp_mode_classique(self):
        self.db.set_setting("bounty_amount", 100)
        a = self.db.add_player("A", "CPC")
        tmp = self.db.add_player("Visiteur", "TMP")
        self.db.eliminate_player(tmp, eliminated_by_id=a)
        bounty_rows = self.db.get_bounty_bonuses()
        names = {r["name"] for r in bounty_rows}
        self.assertNotIn("Visiteur", names)

    def test_bounty_bonuses_exclut_entierement_les_tmp_mode_pko(self):
        self.db.set_setting("pko_mode", 1)
        self.db.set_setting("bounty_amount", 100)
        a = self.db.add_player("A", "CPC")
        tmp = self.db.add_player("Visiteur", "TMP")
        self.db.eliminate_player(tmp, eliminated_by_id=a)
        bounty_rows = self.db.get_bounty_bonuses()
        names = {r["name"] for r in bounty_rows}
        self.assertNotIn("Visiteur", names)

    def test_primes_summary_tmp_tout_a_zero_mais_reste_affiche(self):
        """get_primes_summary (onglet Primes) : un TMP continue
        d'apparaître (trace de sa participation), mais TOUT à 0 —
        aucun cas particulier nécessaire dans cette fonction elle-même,
        résultat naturel des 4 get_*_bonuses déjà corrigées."""
        self.db.set_setting("ranking_formula", database.RANKING_FORMULA_CURRENT)
        self.db.set_setting("bounty_amount", 100)
        a = self.db.add_player("A", "CPC")
        tmp = self.db.add_player("Visiteur", "TMP")
        self.db.eliminate_player(tmp, eliminated_by_id=a)
        summary = {r["name"]: r for r in self.db.get_primes_summary()}
        v = summary["Visiteur"]
        self.assertIsNone(v["rang"])
        self.assertEqual(v["presence"], 0)
        self.assertEqual(v["assiduite"], 0)
        self.assertEqual(v["cl_montant"], 0)
        self.assertEqual(v["bo_montant"], 0)
        self.assertEqual(v["total"], 0)

    def test_primes_summary_tri_par_defaut_tmp_toujours_apres_meme_a_egalite(self):
        """Demande explicite : "les joueurs réguliers doivent toujours
        apparaître avant les TMP dans le tri par défaut, même lorsqu'un
        joueur régulier possède zéro point" — aucun réglage de primes
        activé ici (tous les montants à 0, y compris pour les réguliers),
        pour forcer l'égalité totale."""
        a = self.db.add_player("ZZZ_DernierAlpha", "CPC")  # 0 point, trié dernier par nom si pas de groupe
        tmp = self.db.add_player("AAA_PremierAlpha_TMP", "TMP")  # 0 point aussi, nom alphabétiquement 1er
        b = self.db.add_player("MMM_Milieu", "CPC")
        rows = self.db.get_primes_summary()
        names = [r["name"] for r in rows]
        self.assertEqual(names[-1], "AAA_PremierAlpha_TMP")
        self.assertLess(names.index("ZZZ_DernierAlpha"), names.index("AAA_PremierAlpha_TMP"))
        self.assertLess(names.index("MMM_Milieu"), names.index("AAA_PremierAlpha_TMP"))

    def test_primes_summary_tri_par_rang_inchange_non_regression(self):
        """Non-régression explicite : le tri explicite par 'rang' reste
        INCHANGÉ (il reléguait déjà les TMP en fin de liste via rang=None,
        avant même ce correctif) — jamais retouché."""
        self.db.set_setting("ranking_formula", database.RANKING_FORMULA_CURRENT)
        a = self.db.add_player("A", "CPC")
        tmp = self.db.add_player("Visiteur", "TMP")
        b = self.db.add_player("B", "CPC")
        self.db.eliminate_player(b, eliminated_by_id=a)
        rows = self.db.get_primes_summary(sort_column="rang", ascending=True)
        names = [r["name"] for r in rows]
        self.assertEqual(names[-1], "Visiteur")


# =======================================================================
# 4. players_rows : "rang" (réel, inchangé) vs "official_rang" (nouveau)
# =======================================================================
class PlayersRowsOfficialRangTest(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmpdir, ignore_errors=True)
        self.db = database.Database(os.path.join(self.tmpdir, "A.tournoi"))
        self.addCleanup(self.db.close)

    def test_rang_reel_inclut_le_tmp_official_rang_lexclut(self):
        pierre = self.db.add_player("Pierre", "CPC")
        visiteur = self.db.add_player("Visiteur", "TMP")
        jacques = self.db.add_player("Jacques", "CPC")
        michel = self.db.add_player("Michel", "CPC")
        self.db.eliminate_player(michel, eliminated_by_id=jacques)
        self.db.eliminate_player(jacques, eliminated_by_id=visiteur)
        self.db.eliminate_player(visiteur, eliminated_by_id=pierre)

        rows = {r["name"]: r for r in self.db.players_rows()}
        # "rang" (effectif RÉEL) : ordre brut d'élimination, TMP inclus.
        self.assertEqual(rows["Pierre"]["rang"], 1)
        self.assertEqual(rows["Visiteur"]["rang"], 2)
        self.assertEqual(rows["Jacques"]["rang"], 3)
        self.assertEqual(rows["Michel"]["rang"], 4)
        # "official_rang" (effectif OFFICIEL) : TMP absent (None),
        # réguliers remontés d'un cran chacun après lui.
        self.assertEqual(rows["Pierre"]["official_rang"], 1)
        self.assertIsNone(rows["Visiteur"]["official_rang"])
        self.assertEqual(rows["Jacques"]["official_rang"], 2)
        self.assertEqual(rows["Michel"]["official_rang"], 3)

    def test_joueur_encore_actif_les_deux_rangs_restent_none(self):
        self.db.add_player("Pierre", "CPC")
        self.db.add_player("Jacques", "CPC")
        rows = {r["name"]: r for r in self.db.players_rows()}
        self.assertIsNone(rows["Pierre"]["rang"])
        self.assertIsNone(rows["Pierre"]["official_rang"])


# =======================================================================
# 5. build_period_summary : vainqueur officiel, wins, best_place, played
# =======================================================================
class BuildPeriodSummaryTmpTest(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmpdir, ignore_errors=True)

    def _make_tournament(self, filename, setup_fn):
        path = os.path.join(self.tmpdir, filename)
        db = database.Database(path)
        db.set_setting("primes_enabled", 1)
        db.set_setting("ranking_formula", database.RANKING_FORMULA_CURRENT)
        setup_fn(db)
        db.set_setting("tournament_end_epoch", 1)
        db.close()
        return path

    def test_vainqueur_officiel_quand_un_tmp_gagne_reellement(self):
        def setup(db):
            pierre = db.add_player("Pierre", "CPC")
            visiteur = db.add_player("Visiteur", "TMP")
            jacques = db.add_player("Jacques", "CPC")
            db.eliminate_player(jacques, eliminated_by_id=pierre)
            db.eliminate_player(pierre, eliminated_by_id=visiteur)

        self._make_tournament("To150126.tournoi", setup)
        summary = database.build_period_summary(self.tmpdir)
        self.assertEqual(summary["tournaments"][0]["winner"], "Pierre")

    def test_wins_et_best_place_jamais_pour_un_tmp_meme_vainqueur_reel(self):
        def setup(db):
            pierre = db.add_player("Pierre", "CPC")
            visiteur = db.add_player("Visiteur", "TMP")
            jacques = db.add_player("Jacques", "CPC")
            db.eliminate_player(jacques, eliminated_by_id=pierre)
            db.eliminate_player(pierre, eliminated_by_id=visiteur)

        self._make_tournament("To150126.tournoi", setup)
        summary = database.build_period_summary(self.tmpdir)
        players = {p["name"]: p for p in summary["players"]}
        self.assertEqual(players["Pierre"]["wins"], 1)
        self.assertEqual(players["Pierre"]["best_place"], 1)
        if "Visiteur" in players:
            self.assertEqual(players["Visiteur"]["wins"], 0)
            self.assertIsNone(players["Visiteur"]["best_place"])

    def test_tournaments_played_exclut_les_tmp(self):
        def setup(db):
            pierre = db.add_player("Pierre", "CPC")
            db.add_player("Visiteur", "TMP")
            db.eliminate_player(pierre, eliminated_by_id=None)

        self._make_tournament("To150126.tournoi", setup)
        summary = database.build_period_summary(self.tmpdir)
        players = {p["name"]: p for p in summary["players"]}
        self.assertEqual(players["Pierre"]["tournaments_played"], 1)
        if "Visiteur" in players:
            self.assertEqual(players["Visiteur"]["tournaments_played"], 0)

    def test_entries_reste_leffectif_reel_tmp_compris(self):
        """"entries" (effectif RÉEL, buy-ins) ne doit JAMAIS être réduit
        par ce correctif — distinct de l'effectif OFFICIEL."""
        def setup(db):
            pierre = db.add_player("Pierre", "CPC")
            db.add_player("Visiteur", "TMP")
            db.eliminate_player(pierre, eliminated_by_id=None)

        self._make_tournament("To150126.tournoi", setup)
        summary = database.build_period_summary(self.tmpdir)
        self.assertEqual(summary["tournaments"][0]["entries"], 2)

    def test_sans_aucun_tmp_comportement_strictement_inchange(self):
        """Non-régression centrale : un tournoi SANS aucun joueur "TMP"
        produit des résultats identiques à avant ce correctif."""
        def setup(db):
            pierre = db.add_player("Pierre", "CPC")
            jacques = db.add_player("Jacques", "CPC")
            db.eliminate_player(jacques, eliminated_by_id=pierre)

        self._make_tournament("To150126.tournoi", setup)
        summary = database.build_period_summary(self.tmpdir)
        self.assertEqual(summary["tournaments"][0]["winner"], "Pierre")
        players = {p["name"]: p for p in summary["players"]}
        self.assertEqual(players["Pierre"]["wins"], 1)
        self.assertEqual(players["Pierre"]["tournaments_played"], 1)
        self.assertEqual(players["Jacques"]["tournaments_played"], 1)

    def test_forfait_sans_tmp_najamais_de_vainqueur_invente(self):
        """PIÈGE RÉEL détecté en test (voir OfficialStandingsTest::
        test_trou_sans_rapport_avec_un_tmp_najamais_promu_a_1 et tests/
        test_period_summary_stats_tab.py::
        ForfaitNeCompteJamaisCommeTournoiJoueTest, déjà existant) : un
        tournoi où tout le monde finit éliminé OU forfait (0 actif, AUCUN
        TMP impliqué) ne doit jamais désigner un "vainqueur" — même si le
        dernier joueur éliminé se retrouve, par la renumérotation des
        places existante (voir Database.withdraw_player), avec une place
        brute de 1."""
        def setup(db):
            a = db.add_player("Alice", "CPC")
            b = db.add_player("Bob", "CPC")
            db.eliminate_player(b, eliminated_by_id=a)
            db.withdraw_player(a)

        self._make_tournament("To150126.tournoi", setup)
        summary = database.build_period_summary(self.tmpdir)
        self.assertEqual(summary["tournaments"][0]["winner"], "-")


# =======================================================================
# 6. Rétroactivité : un ancien tournoi (sans TMP) reste inchangé
# =======================================================================
class RetroactivityNonRegressionTest(unittest.TestCase):
    """Simule un tournoi "ancien" (créé avant ce correctif, donc sans
    aucune notion de TMP) : aucune colonne n'a été ajoutée au schéma, ce
    tournoi continue de fonctionner exactement comme avant, sans la
    moindre migration ni changement de résultat."""

    def test_club_vide_jamais_confondu_avec_tmp(self):
        tmpdir = tempfile.mkdtemp()
        try:
            db = database.Database(os.path.join(tmpdir, "A.tournoi"))
            db.set_setting("primes_enabled", 1)
            db.set_setting("ranking_formula", database.RANKING_FORMULA_CURRENT)
            db.set_setting("attendance_bonus_points", 10)
            # Club jamais renseigné (comportement des anciens tournois,
            # avant même l'existence du champ club) : chaîne vide.
            a = db.add_player("A", "")
            b = db.add_player("B", "")
            db.eliminate_player(b, eliminated_by_id=a)
            ranking = {r["name"]: r["place"] for r in db.get_ranking_bonuses()}
            self.assertEqual(ranking, {"A": 1, "B": 2})
            presence = db.get_presence_bonuses()
            self.assertEqual(presence["A"], 10)
            self.assertEqual(presence["B"], 10)
            db.close()
        finally:
            shutil.rmtree(tmpdir, ignore_errors=True)


# =======================================================================
# 7. Frontière de périmètre : les gains réels (Résultats/Gains) restent
#    INCHANGÉS, basés sur l'ordre RÉEL — hors périmètre de ce correctif.
# =======================================================================
class ResultsRowsScopeBoundaryTest(unittest.TestCase):
    """Verrouille explicitement que _results_rows/get_payouts_amounts
    (onglet "Résultats", gains en argent réel) ne sont PAS touchés par ce
    correctif — décision de périmètre explicite (voir le compte-rendu),
    aucune règle de l'utilisateur ne portant sur les gains réels."""

    def test_results_rows_utilise_toujours_la_place_brute(self):
        import inspect
        tree = ast.parse(inspect.getsource(database), filename="database.py")
        func = next(
            n for n in ast.walk(tree)
            if isinstance(n, ast.FunctionDef) and n.name == "_results_rows"
        )
        source = ast.get_source_segment(inspect.getsource(database), func)
        self.assertNotIn("official_standings", source)
        self.assertNotIn("is_tmp_club", source)
        self.assertIn('p["place"]', source)


# =======================================================================
# 8. Câblage main.py : l'onglet Classement affiche bien le rang OFFICIEL
# =======================================================================
class _MainSourceHelper(unittest.TestCase):
    def setUp(self):
        self.source = _main_source()
        self.tree = ast.parse(self.source, filename="main.py")

    def _find_method(self, class_name, method_name):
        cls_node = next(
            n for n in ast.walk(self.tree)
            if isinstance(n, ast.ClassDef) and n.name == class_name
        )
        return next(
            n for n in ast.walk(cls_node)
            if isinstance(n, ast.FunctionDef) and n.name == method_name
        )

    def _source_of(self, node):
        return ast.get_source_segment(self.source, node) or ""


class ClassementTabStructuralTest(_MainSourceHelper):
    """main.py construit une vraie fenêtre Tk complète pour l'onglet
    Classement (coûteux/fragile à instancier ici, même limite déjà
    documentée ailleurs dans cette suite) : vérification STRUCTURELLE du
    code source, même convention déjà établie (voir tests/test_backup_
    restore.py::RestoreMultiFolderTransparencyTest et consorts).

    Précision du 2026-10-15 (après une première version qui excluait à
    tort les TMP de l'onglet Classement) : voir ClassementRowsTest et
    ClassementExportTest ci-dessous pour le comportement RÉEL (lignes
    renvoyées), cette classe-ci ne verrouille que le CÂBLAGE (quelle
    fonction/quel champ est appelé)."""

    def test_classement_rows_utilise_database_classement_rows(self):
        func = self._find_method("App", "_classement_rows")
        seg = self._source_of(func)
        self.assertIn("self.db.classement_rows()", seg)
        # Jamais retombé sur players_rows/official_rang directement :
        # toute la logique régulier/TMP vit désormais UNIQUEMENT dans
        # Database.classement_rows (jamais dupliquée ici).
        self.assertNotIn("players_rows", seg)
        self.assertNotIn("official_rang", seg)

    def test_affichage_du_tableau_classement_utilise_rang_deja_resolu(self):
        func = self._find_method("App", "_refresh_payouts_tab")
        seg = self._source_of(func)
        self.assertIn('r["rang"]', seg)

    def test_export_classement_utilise_les_fonctions_dediees(self):
        """Demande explicite : "sans modifier le comportement des
        exports génériques de l'onglet Joueurs" — ClassementExportDialog
        doit utiliser les 3 fonctions DÉDIÉES, jamais export_players_*."""
        func = self._find_method("ClassementExportDialog", "_do_export")
        seg = self._source_of(func)
        self.assertIn("self.db.export_classement_xlsx(", seg)
        self.assertIn("self.db.export_classement_pdf(", seg)
        self.assertIn("self.db.export_classement_csv(", seg)
        # Un commentaire explique pourquoi (voir docstring de la méthode)
        # mais aucun APPEL réel aux fonctions génériques Joueurs.
        self.assertNotIn("self.db.export_players_xlsx(", seg)
        self.assertNotIn("self.db.export_players_pdf(", seg)
        self.assertNotIn("self.db.export_players_csv(", seg)


class PlayersTabUnaffectedStructuralTest(_MainSourceHelper):
    """Demande explicite : "Onglet Joueurs — ne rien changer à son
    fonctionnement actuel" — verrouille que _build_players_tab/
    players_rows/export_players_* restent entièrement indépendants de
    classement_rows/is_tmp_club côté affichage (le filtrage TMP des
    POINTS, lui, reste correct puisqu'il vit dans database.py, jamais
    dans l'onglet Joueurs lui-même, qui n'affiche aucun point)."""

    def test_export_players_jamais_modifiees_par_le_groupe_tmp(self):
        import inspect
        import database as db_module
        src = inspect.getsource(db_module)
        tree = ast.parse(src, filename="database.py")
        for name in ("export_players_csv", "export_players_xlsx", "export_players_pdf"):
            func = next(
                n for n in ast.walk(tree)
                if isinstance(n, ast.FunctionDef) and n.name == name
            )
            seg = ast.get_source_segment(src, func)
            self.assertNotIn("classement_rows", seg)
            self.assertNotIn("is_tmp_club", seg)
            self.assertIn("players_rows", seg)


class ClassementRowsTest(unittest.TestCase):
    """Database.classement_rows — comportement réel, cas particuliers
    explicitement demandés : aucun TMP, plusieurs TMP, TMP vainqueur
    réel, joueur régulier à zéro point, joueurs encore actifs, forfaits."""

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmpdir, ignore_errors=True)
        self.db = database.Database(os.path.join(self.tmpdir, "A.tournoi"))
        self.addCleanup(self.db.close)

    def test_exemple_exact_reguliers_dabord_puis_tmp_hors_classement(self):
        pierre = self.db.add_player("Pierre", "CPC")
        visiteur = self.db.add_player("Visiteur", "TMP")
        jacques = self.db.add_player("Jacques", "CPC")
        michel = self.db.add_player("Michel", "CPC")
        self.db.eliminate_player(michel, eliminated_by_id=jacques)
        self.db.eliminate_player(jacques, eliminated_by_id=visiteur)
        self.db.eliminate_player(visiteur, eliminated_by_id=pierre)
        rows = self.db.classement_rows()
        self.assertEqual(
            [(r["name"], r["rang"]) for r in rows],
            [("Pierre", 1), ("Jacques", 2), ("Michel", 3), ("Visiteur", "Hors classement")],
        )

    def test_aucun_tmp_comportement_normal(self):
        pierre = self.db.add_player("Pierre", "CPC")
        jacques = self.db.add_player("Jacques", "CPC")
        self.db.eliminate_player(jacques, eliminated_by_id=pierre)
        rows = self.db.classement_rows()
        self.assertEqual([(r["name"], r["rang"]) for r in rows], [("Pierre", 1), ("Jacques", 2)])

    def test_plusieurs_tmp_tous_apres_les_reguliers_jamais_entremeles(self):
        pierre = self.db.add_player("Pierre", "CPC")
        v1 = self.db.add_player("Visiteur1", "TMP")
        v2 = self.db.add_player("Visiteur2", "TMP")
        jacques = self.db.add_player("Jacques", "CPC")
        self.db.eliminate_player(v2, eliminated_by_id=jacques)
        self.db.eliminate_player(jacques, eliminated_by_id=v1)
        self.db.eliminate_player(v1, eliminated_by_id=pierre)
        # Ordre réel d'élimination : v2 (place brute 4), jacques (3),
        # v1 (2) ; Pierre reste seul actif (1). Les TMP gardent, entre
        # eux, l'ordre de leur place brute (Visiteur1 éliminé après
        # Visiteur2, donc sa place brute 2 est "meilleure" que la 4).
        rows = self.db.classement_rows()
        names = [r["name"] for r in rows]
        rangs = [r["rang"] for r in rows]
        self.assertEqual(names, ["Pierre", "Jacques", "Visiteur1", "Visiteur2"])
        self.assertEqual(rangs, [1, 2, "Hors classement", "Hors classement"])

    def test_tmp_vainqueur_reel_meilleur_regulier_devient_rang_1(self):
        pierre = self.db.add_player("Pierre", "CPC")
        visiteur = self.db.add_player("Visiteur", "TMP")
        jacques = self.db.add_player("Jacques", "CPC")
        self.db.eliminate_player(jacques, eliminated_by_id=pierre)
        self.db.eliminate_player(pierre, eliminated_by_id=visiteur)
        # Visiteur (TMP) reste seul actif -> vainqueur RÉEL.
        rows = self.db.classement_rows()
        self.assertEqual(
            [(r["name"], r["rang"]) for r in rows],
            [("Pierre", 1), ("Jacques", 2), ("Visiteur", "Hors classement")],
        )

    def test_regulier_a_zero_point_reste_classe_normalement(self):
        """Un régulier classé n'a ici ni point ni total (classement_rows
        ne montre pas de points, voir sa docstring) — seul son RANG
        officiel compte, jamais affecté par le fait qu'il ait 0 point
        ailleurs (onglet Primes)."""
        self.db.set_setting("primes_enabled", 0)  # aucun point nulle part
        pierre = self.db.add_player("Pierre", "CPC")
        jacques = self.db.add_player("Jacques", "CPC")
        self.db.eliminate_player(jacques, eliminated_by_id=pierre)
        rows = self.db.classement_rows()
        self.assertEqual([(r["name"], r["rang"]) for r in rows], [("Pierre", 1), ("Jacques", 2)])

    def test_joueurs_encore_actifs_jamais_dans_classement_rows(self):
        self.db.add_player("Pierre", "CPC")
        self.db.add_player("Visiteur", "TMP")
        self.assertEqual(self.db.classement_rows(), [])

    def test_forfait_jamais_dans_classement_rows_tmp_ou_non(self):
        a = self.db.add_player("Alice", "CPC")
        b = self.db.add_player("Bob", "TMP")
        self.db.withdraw_player(a)
        self.db.withdraw_player(b)
        self.assertEqual(self.db.classement_rows(), [])

    def test_ne_renvoie_jamais_les_colonnes_de_players_rows_sans_rapport(self):
        pierre = self.db.add_player("Pierre", "CPC")
        self.db.eliminate_player(pierre, eliminated_by_id=None)
        rows = self.db.classement_rows()
        self.assertEqual(
            set(rows[0].keys()), {"name", "club", "rang", "elim_time", "elim_round", "eliminated_by"},
        )


class ClassementExportTest(unittest.TestCase):
    """Database.export_classement_csv/xlsx/pdf — mêmes règles que
    classement_rows, jamais mélangées, jamais via export_players_*."""

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmpdir, ignore_errors=True)
        self.db = database.Database(os.path.join(self.tmpdir, "A.tournoi"))
        self.addCleanup(self.db.close)
        pierre = self.db.add_player("Pierre", "CPC")
        visiteur = self.db.add_player("Visiteur", "TMP")
        jacques = self.db.add_player("Jacques", "CPC")
        self.db.eliminate_player(jacques, eliminated_by_id=pierre)
        self.db.eliminate_player(pierre, eliminated_by_id=visiteur)
        self.db.eliminate_player(visiteur, eliminated_by_id=None)

    def test_csv_reguliers_dabord_tmp_ensuite_hors_classement(self):
        path = os.path.join(self.tmpdir, "out.csv")
        self.db.export_classement_csv(path)
        with open(path, encoding="utf-8-sig") as f:
            lines = [ln.strip() for ln in f.read().splitlines() if ln.strip()]
        # En-tête + 3 joueurs.
        self.assertEqual(len(lines), 4)
        self.assertTrue(lines[1].startswith("Pierre;"))
        self.assertTrue(lines[-1].startswith("Visiteur;"))
        self.assertIn("Hors classement", lines[-1])

    def test_xlsx_reguliers_dabord_tmp_ensuite_hors_classement(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl indisponible dans cet environnement")
        path = os.path.join(self.tmpdir, "out.xlsx")
        self.db.export_classement_xlsx(path)
        wb = load_workbook(path)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        # Titre + ligne vide + en-tête + 3 joueurs (voir export_classement_xlsx).
        data_rows = rows[3:]
        self.assertEqual(data_rows[0][0], "Pierre")
        self.assertEqual(data_rows[-1][0], "Visiteur")
        self.assertEqual(data_rows[-1][1], "Hors classement")

    def test_pdf_sexporte_sans_erreur(self):
        try:
            import fpdf  # noqa: F401
        except ImportError:
            self.skipTest("fpdf2 indisponible dans cet environnement")
        path = os.path.join(self.tmpdir, "out.pdf")
        self.db.export_classement_pdf(path)
        self.assertTrue(os.path.isfile(path))
        self.assertGreater(os.path.getsize(path), 0)

    def test_export_vide_si_aucun_joueur_classe(self):
        """Deux joueurs encore actifs (ni l'un ni l'autre "dernier
        survivant" — voir Database.players_rows : un SEUL actif restant
        obtient automatiquement le rang 1) : aucune ligne classée."""
        tmpdir2 = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, tmpdir2, ignore_errors=True)
        db2 = database.Database(os.path.join(tmpdir2, "B.tournoi"))
        self.addCleanup(db2.close)
        db2.add_player("EncoreActif1", "CPC")
        db2.add_player("EncoreActif2", "CPC")
        path = os.path.join(tmpdir2, "vide.csv")
        db2.export_classement_csv(path)
        with open(path, encoding="utf-8-sig") as f:
            lines = [ln.strip() for ln in f.read().splitlines() if ln.strip()]
        self.assertEqual(len(lines), 1)  # en-tête seul


if __name__ == "__main__":
    unittest.main()
